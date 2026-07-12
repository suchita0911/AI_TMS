"""Render report data to Excel and PDF byte streams."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    Paragraph,
)

_HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_excel(sections: dict[str, tuple[list[str], list[list]]]) -> bytes:
    """sections: {sheet_name: (headers, rows)}."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sections.items():
        ws = wb.create_sheet(title=name[:31])
        ws.append(headers)
        for c in ws[1]:
            c.fill = _HEADER_FILL
            c.font = _HEADER_FONT
        for row in rows:
            ws.append(row)
        for i, h in enumerate(headers, start=1):
            width = max(len(str(h)), *(len(str(r[i - 1])) for r in rows)) if rows else len(str(h))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width + 4, 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(title: str, sections: dict[str, tuple[list[str], list[list]]]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=title,
                            leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 8)]

    for name, (headers, rows) in sections.items():
        elements.append(Paragraph(name, styles["Heading2"]))
        data = [headers] + (rows if rows else [["—"] * len(headers)])
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 12))

    doc.build(elements)
    return buf.getvalue()
